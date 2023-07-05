from docx import Document
from utils import fix_paragraphs, merge_runs, get_attribute, is_metadata
from segment import segment2idx, SegmentClassifier, SegmentType
import csv


file_path = "mevzuat-raw/dmk.docx"
document = Document(file_path)
output_file = open("output-dmk.csv", "w", encoding="utf-8")
csv_writer = csv.writer(output_file)

fixed_paragraphs = fix_paragraphs(document.paragraphs)

#Create an auxiliary list to hold all runs in the document
flattened_runs = [] 
for paragraph in fixed_paragraphs:
    flattened_runs.extend(merge_runs(paragraph.runs))

classifier = SegmentClassifier(flattened_runs)

#Handle all classes except footnotes and tables (and infrequent classes)
for parag_idx, paragraph in enumerate(fixed_paragraphs):
    fixed_runs = merge_runs(paragraph.runs)
    for run_idx, run in enumerate(fixed_runs):
        segment_class = classifier.classify_next(run)
        if not segment_class:
            continue
        
        csv_writer.writerow([parag_idx, run_idx, run.text, segment_class.name, get_attribute(run, "bold"), get_attribute(run, "italic"), get_attribute(run, "alignment")])
    
#Handle the footnotes from the original documant.paragraphs object
document = Document(file_path)
for paragraph in document.paragraphs:
    for run in paragraph.runs:
        for footnote_ref in run.footnotes:
            for foot_parag in footnote_ref.footnote.paragraphs:
                if foot_parag.text:
                    csv_writer.writerow([parag_idx, run_idx, foot_parag.text, segment2idx["FOOTNOTE"]])


document = Document(file_path)
fixed_paragraphs = fix_paragraphs(document.paragraphs)
for parag_idx, paragraph in enumerate(fixed_paragraphs):
    for run_idx, run in enumerate(paragraph.runs):
        if is_metadata(run):
            csv_writer.writerow([parag_idx, run_idx, run.text, SegmentType.METADATA, get_attribute(run, "bold"), get_attribute(run, "italic"), get_attribute(run, "alignment")])
            
output_file.close()